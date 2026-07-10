#!/usr/bin/env python3
"""Convert the Cornerstone assignment-pp_cv1 form into a generator template.

Replaces bracket placeholders with Jinja fields and writes
templates/assignment-of-personal-property.docx for the
assignment-of-personal-property package.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "transfers" / "cornerstone" / "assignment-pp_cv1.docx"
TARGET = ROOT / "templates" / "assignment-of-personal-property.docx"

TRUSTEE_GRANTEE_LINE = (
    "{{ fiduciaries.trustee.primary.full_name | name }}"
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    " and {{ fiduciaries.trustee.co_primary.full_name | name }}"
    "{% endif %} as Trustee"
    "{% if fiduciaries.trustee.co_primary.full_name %}s{% endif %}, or their successors in interest, "
    "of the {{ trust.display_name }} dated {{ trust.date }}, and any amendments thereto"
)

TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "[Trustee] and [Co-Trustee] as Trustees, or their successors in interest, of the [TRUST NAME] dated [DocDate], and any amendments thereto",
        TRUSTEE_GRANTEE_LINE,
    ),
    ("[ADDRESS]", "{{ client.street_address }}, {{ client.city }}, {{ client.state }}"),
    ("[SIGNING COUNTY]", "{{ execution.signing_county | name }}"),
    ("[Notary Commission]", "{{ execution.notary_commission }}"),
    ("[client heshe]", "{{ client.pronoun_subject }}"),
    ("[TRUST NAME]", "{{ trust.display_name }}"),
    ("[CLIENT]", "{{ client.full_name | name }}"),
    ("[Trustee]", "{{ fiduciaries.trustee.primary.full_name | name }}"),
    ("[Co-Trustee]", "{{ fiduciaries.trustee.co_primary.full_name | name }}"),
    ("[DocDate]", "{{ execution.doc_date }}"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
