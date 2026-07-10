#!/usr/bin/env python3
"""Convert the Cornerstone certificate-trust_cv1 form into a generator template.

Replaces bracket placeholders with Jinja fields and writes
templates/certificate-of-trust.docx for the certificate-of-trust package.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "transfers" / "cornerstone" / "certificate-trust_cv1.docx"
TARGET = ROOT / "templates" / "certificate-of-trust.docx"

TRUSTEE_TITLE_LINE = (
    "{{ fiduciaries.trustee.primary.full_name | name }}"
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    " and {{ fiduciaries.trustee.co_primary.full_name | name }}"
    "{% endif %} as Trustees, or their successors in interest of the {{ trust.display_name }} "
    "dated {{ trust.date }}, and any amendments thereto"
)

TRUSTEE_ADDRESS_LINE = (
    "The current trustees are {{ fiduciaries.trustee.primary.full_name | name }}"
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    " and {{ fiduciaries.trustee.co_primary.full_name | name }}"
    "{% endif %} of {{ client.street_address }}, {{ client.city }}, {{ client.state }}."
)

MULTI_TRUSTEE_PARAGRAPH = (
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    "The trust does have multiple trustees.  The signature of one of the trustees is required to exercise the powers of the trust."
    "{% else %}"
    "The trust has a single trustee."
    "{% endif %}"
)

NOTARY_GRANTOR = (
    "On {{ execution.ordinal_doc_date }}, before me, the undersigned notary public, personally appeared "
    "{{ client.full_name | name }}, ___ personally known to me, or ___ proved to me through satisfactory "
    "evidence of identification, which was a government-issued photo identification to be the person whose "
    "name is signed on the preceding or attached document, and who acknowledged that "
    "{{ client.pronoun_subject }} signed it voluntarily for its stated purpose."
)

NOTARY_CO_TRUSTEE = (
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    "On {{ execution.ordinal_doc_date }}, before me, the undersigned notary public, personally appeared "
    "{{ fiduciaries.trustee.co_primary.full_name | name }}, ___ personally known to me, or ___ proved to me "
    "through satisfactory evidence of identification, which was a government-issued photo identification "
    "to be the person whose name is signed on the preceding or attached document, and who acknowledged that "
    "{{ fiduciaries.trustee.co_primary.pronoun_subject | default(client.pronoun_subject) }} signed it "
    "voluntarily for its stated purpose."
    "{% endif %}"
)

TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "[INITIAL TRUSTEE] and [CO-TRUSTEE] as Trustees, or their successors in interest of the [TRUST NAME] dated [DocDate], and any amendments thereto",
        TRUSTEE_TITLE_LINE,
    ),
    (
        "The current trustees are [INITIAL TRUSTEE] and [CO-TRUSTEE] of [GRANTOR ADDRESS].",
        TRUSTEE_ADDRESS_LINE,
    ),
    (
        "The trust does have multiple trustees.  The signature of one of the trustees is required to exercise the powers of the trust.",
        MULTI_TRUSTEE_PARAGRAPH,
    ),
    ("[GRANTOR ADDRESS]", "{{ client.street_address }}, {{ client.city }}, {{ client.state }}"),
    ("[SIGNING COUNTY]", "{{ execution.signing_county | name }}"),
    ("[NotaryExpiration]", "{{ execution.notary_commission }}"),
    ("[cotrustee he/she]", "{{ fiduciaries.trustee.co_primary.pronoun_subject | default(client.pronoun_subject) }}"),
    ("[grantor he/she]", "{{ client.pronoun_subject }}"),
    ("[INITIAL TRUSTEE]", "{{ fiduciaries.trustee.primary.full_name | name }}"),
    ("[CO-TRUSTEE]", "{{ fiduciaries.trustee.co_primary.full_name | name }}"),
    ("[TRUST NAME]", "{{ trust.display_name }}"),
    ("[GRANTOR]", "{{ client.full_name | name }}"),
    ("[DocDate]", "{{ execution.ordinal_doc_date }}"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def apply_paragraph_fixes(document: Document) -> None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph_text(paragraph)

        if text.startswith("On [DocDate], before me, the undersigned notary public, personally appeared [GRANTOR]"):
            replace_paragraph_text(paragraph, NOTARY_GRANTOR)
            continue

        if text.startswith("On [DocDate], before me, the undersigned notary public, personally appeared [COTRUSTEE]"):
            replace_paragraph_text(paragraph, NOTARY_CO_TRUSTEE)
            continue

        if "[COTRUSTEE]" in text and "notary public" in text:
            replace_paragraph_text(paragraph, NOTARY_CO_TRUSTEE)
            continue

        if text == "[GRANTOR], Trustee":
            replace_paragraph_text(
                paragraph,
                "{{ client.full_name | name }}, Trustee",
            )
            continue

        if text == "[CO-TRUSTEE], Trustee":
            replace_paragraph_text(
                paragraph,
                "{% if fiduciaries.trustee.co_primary.full_name %}"
                "{{ fiduciaries.trustee.co_primary.full_name | name }}, Trustee"
                "{% endif %}",
            )
            continue


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    apply_paragraph_fixes(document)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
