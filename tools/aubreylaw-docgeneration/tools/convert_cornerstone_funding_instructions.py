#!/usr/bin/env python3
"""Convert the Cornerstone funding-instructions_cv1 form into a generator template.

Replaces ALL-CAPS placeholders with Jinja fields and writes
templates/funding-instructions.docx for the funding-instructions package.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "transfers" / "cornerstone" / "funding-instructions_cv1.docx"
TARGET = ROOT / "templates" / "funding-instructions.docx"

SUBTITLE_LINE = (
    "{% if has_spouse %}"
    "The {{ client.full_name | name }} and The {{ spouse.full_name | name }}"
    "{% else %}"
    "The {{ trust.display_name }}"
    "{% endif %}"
)

TITLING_FORMAT_LINE = (
    "{{ fiduciaries.trustee.primary.full_name | name }}"
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    " and {{ fiduciaries.trustee.co_primary.full_name | name }}"
    "{% endif %}, Trustees, or their successors in interest, of the {{ trust.name | name }} "
    "dated {{ trust.date }}, and any amendments thereto."
)

TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    ("The TRUST NAME and The TRUST NAME", SUBTITLE_LINE),
    (
        "TRUSTEE and COTRUSTEE, Trustees, or their successors in interest, of the TRUST NAME dated _____________, 20___, and any amendments thereto.",
        TITLING_FORMAT_LINE,
    ),
    ("TRUST NAME", "{{ trust.display_name }}"),
    ("TRUSTEE", "{{ fiduciaries.trustee.primary.full_name | name }}"),
    ("COTRUSTEE", "{{ fiduciaries.trustee.co_primary.full_name | name }}"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
