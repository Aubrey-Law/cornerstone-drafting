#!/usr/bin/env python3
"""Convert the Cornerstone homestead-trustee_cv1 form into a generator template.

Replaces bracket placeholders with Jinja fields and writes
templates/homestead-trustee.docx for the homestead-trustee package.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "transfers" / "cornerstone" / "homestead-trustee_cv1.docx"
TARGET = ROOT / "templates" / "homestead-trustee.docx"

TRUSTEE_NAMES = (
    "{{ fiduciaries.trustee.primary.full_name | name }}"
    "{% if fiduciaries.trustee.co_primary.full_name %}"
    " and {{ fiduciaries.trustee.co_primary.full_name | name }}"
    "{% endif %}"
)

OPENING_DECLARATION = (
    "I/We, "
    + TRUSTEE_NAMES
    + ", as Trustee(s) of the {{ trust.name | name }}, dated {{ trust.date }}, "
    "recorded in Book {{ homestead.trust_deed_book }}, Page {{ homestead.trust_deed_page }}, "
    "hereby declare homestead pursuant to M.G.L. c. 188 and state that I/we own the home described below "
    "and which the beneficiaries listed herein occupy or intend to occupy as "
    "{{ homestead.beneficiary_pronoun_possessive }} principal residence."
)

BENEFICIARY_LINE = (
    "The beneficiary/beneficiaries of this trust is/are: {{ homestead.beneficiary_names }}."
)

ELDERLY_LINE = (
    "{{ homestead.beneficiary_names }} is/are elderly (62 years of age or older)."
)

DISABLED_LINE = (
    "{{ homestead.beneficiary_names }} is/are disabled (have a physical or mental impairment that meets "
    "the disability requirements for Supplemental Security Income under 42 U.S.C. 1382c(a)(3)(A) and "
    "42 U.S.C. 1382c(a)(3)(C)). One of the following must be attached: 1) an original or certified copy "
    "of a disability award letter issued to the person by the United States Social Security Administration, "
    "or 2) a letter signed by a physician registered with the board of registration in medicine certifying "
    "that each person meets the disability requirements stated in 42 U.S.C. 1382c(a)(3)(A) and "
    "42 U.S.C. 1382c(a)(3)(C)."
)

SERVICEMEMBER_LINE = (
    "{% if homestead.servicemember_names %}"
    "I/We, {{ homestead.servicemember_names }}, am/are servicemember(s) who may be subject to protection "
    "under the Servicemembers Civil Relief Act, 50 U.S.C. App 533, should I/we be called to active duty."
    "{% endif %}"
)

SPOUSE_OCCUPANT_LINE = (
    "{% if has_spouse %}"
    "{{ homestead.primary_beneficiary_name }} is married to {{ spouse.full_name | name }}, who is not a "
    "co-owner of the home but who occupies or intends to occupy the home as "
    "{{ spouse.pronoun_possessive | default('their') }} principal residence."
    "{% endif %}"
)

PROPERTY_ADDRESS_LINE = (
    "Address: {{ homestead.property_address_line }}, Massachusetts."
)

DEED_RECORDING_LINE = (
    "Deed is recorded in the {{ homestead.deed_registry_district }} Registry of Deeds in Book "
    "{{ homestead.deed_book }}, Page {{ homestead.deed_page }}. Certificate of Title No. "
    "{{ homestead.certificate_number }}, registered in the Land Registration Office, Document No. "
    "{{ homestead.land_court_document_number }}."
)

NOTARY_BLOCK = (
    "On this {{ execution.ordinal_doc_date }} before me, the undersigned notary public, personally appeared "
    + TRUSTEE_NAMES
    + ", proved to me through satisfactory evidence of identification, which was a government-issued photo "
    "identification, to be the person(s) who signed the preceding or attached document in my presence, "
    "and who swore or affirmed to me that the contents of the document are truthful and accurate to the "
    "best of {{ homestead.trustee_pronoun_possessive }} knowledge and belief."
)

TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    ("[Trustee Full Name(s)]", TRUSTEE_NAMES),
    ("[Trust Name]", "{{ trust.display_name }}"),
    ("[Trust Date]", "{{ trust.date }}"),
    ("[Beneficiary Name(s)]", "{{ homestead.beneficiary_names }}"),
    ("[Beneficiary Name]", "{{ homestead.primary_beneficiary_name }}"),
    ("[Spouse Name]", "{{ spouse.full_name | name }}"),
    ("[Servicemember Name(s)]", "{{ homestead.servicemember_names }}"),
    (
        "[Street Number and Name, City/Town]",
        "{{ homestead.property_address_line }}",
    ),
    ("[District/County]", "{{ homestead.deed_registry_district }}"),
    ("[Certificate No.]", "{{ homestead.certificate_number }}"),
    ("[Document No.]", "{{ homestead.land_court_document_number }}"),
    ("[OrdinalDocDate]", "{{ execution.ordinal_doc_date }}"),
    ("[SINGING COUNTY", "{{ execution.signing_county | name }}"),
    ("[SIGNING COUNTY]", "{{ execution.signing_county | name }}"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def apply_paragraph_fixes(document: Document) -> None:
    paragraphs = document.paragraphs
    spouse_line_seen = False
    for paragraph in paragraphs:
        text = paragraph_text(paragraph)

        if text.startswith("I/We, [Trustee Full Name(s)]"):
            replace_paragraph_text(paragraph, OPENING_DECLARATION)
            continue

        if text.startswith("The beneficiary/beneficiaries of this trust"):
            replace_paragraph_text(paragraph, BENEFICIARY_LINE)
            continue

        if "[Beneficiary Name(s)] is/are elderly" in text:
            replace_paragraph_text(paragraph, ELDERLY_LINE)
            continue

        if "[Beneficiary Name(s)] is/are disabled" in text:
            replace_paragraph_text(paragraph, DISABLED_LINE)
            continue

        if text.startswith("I/We, [Servicemember Name(s)]"):
            replace_paragraph_text(paragraph, SERVICEMEMBER_LINE)
            continue

        if "[Beneficiary Name] is married to [Spouse Name]" in text:
            if spouse_line_seen:
                replace_paragraph_text(paragraph, "")
            else:
                replace_paragraph_text(paragraph, SPOUSE_OCCUPANT_LINE)
                spouse_line_seen = True
            continue

        if text.startswith("Address: [Street Number and Name"):
            replace_paragraph_text(paragraph, PROPERTY_ADDRESS_LINE)
            continue

        if text.startswith("Deed is recorded in the [District/County]"):
            replace_paragraph_text(paragraph, DEED_RECORDING_LINE)
            continue

        if text.startswith("On this [OrdinalDocDate] before me"):
            replace_paragraph_text(paragraph, NOTARY_BLOCK)
            continue


def apply_signature_blocks(document: Document) -> None:
    trustee_labels = [p for p in document.paragraphs if paragraph_text(p) == "[TRUSTEE]"]

    if len(trustee_labels) >= 1:
        replace_paragraph_text(
            trustee_labels[0],
            "{{ fiduciaries.trustee.primary.full_name | name }}, Trustee",
        )
    if len(trustee_labels) >= 2:
        replace_paragraph_text(
            trustee_labels[1],
            "{% if fiduciaries.trustee.co_primary.full_name %}"
            "{{ fiduciaries.trustee.co_primary.full_name | name }}, Trustee"
            "{% endif %}",
        )
    for extra in trustee_labels[2:]:
        replace_paragraph_text(extra, "")


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    apply_paragraph_fixes(document)
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    apply_signature_blocks(document)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
