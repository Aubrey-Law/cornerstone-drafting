#!/usr/bin/env python3
"""Convert the Cornerstone Power of Attorney_cv1 form into a generator template.

The source form already contains Jinja conditionals for three opt-in modules:
  - dpoa.include_entity_powers
  - dpoa.include_advanced_trust_gifting
  - dpoa.include_extended_personal_care

This script replaces sample client data with Jinja placeholders and writes
templates/durable-power-of-attorney.docx for the incapacity-documents package.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "incapacity" / "cornerstone" / "Power of Attorney_cv1.docx"
TARGET = ROOT / "templates" / "durable-power-of-attorney.docx"

# Order matters: longer / more specific strings first.
TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "I, [CLIENT FULL NAME] of Needham, Massachusetts, am creating a Durable Power of Attorney",
        "I, {{ client.full_name | name }} of {{ client.city }}, {{ client.state }}, am creating a Durable Power of Attorney",
    ),
    ("[CLIENT FULL NAME]", "{{ client.full_name | name }}"),
    ("[Initial Health Care Proxy]", "{{ fiduciaries.agent.primary.full_name | name }}"),
    ("[Initial Attorney-in-Fact]", "{{ fiduciaries.agent.primary.full_name | name }}"),
    ("[Successor Attorney-in-Fact]", "{{ fiduciaries.agent.successor.full_name | name }}"),
    ("COUNTY OF [SIGNING COUNTY", "COUNTY OF {{ execution.display_signing_county | name }}"),
    ("[DocDate]", "{{ execution.display_doc_date }}"),
    ("[Notary Commission]", "{{ execution.display_notary_commission }}"),
    ("My Commission Expires: November 1, 2030", ""),
]


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph) -> None:
    original = paragraph.text
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
