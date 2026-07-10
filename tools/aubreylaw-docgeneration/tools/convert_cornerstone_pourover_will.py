#!/usr/bin/env python3
"""Convert the Cornerstone pour-over will_cv1 form into a generator template.

The cv1 source is mostly pre-templated with Jinja placeholders and clause includes.
This script removes legacy duplicate PR paragraphs, fixes the PR nomination
conditionals, applies family/disinheritance conditional patches, and writes
templates/pourover-will.docx for the pourover-will package.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "estate" / "cornerstone" / "pourover-will_cv1.docx"
TARGET = ROOT / "templates" / "pourover-will.docx"

# Legacy duplicate paragraphs left in cv1 with bracket placeholders.
DELETE_PARAGRAPH_PREFIXES = (
    "I appoint [Initial Personal Representative]",
    "Should all of the said [CLIENT PR]",
)

# Fixed Cornerstone PR nomination (cv1 legal text with working Jinja conditionals).
PR_NOMINATION_PARAGRAPH = (
    "I nominate {{ fiduciaries.executor.primary.full_name | name }}"
    "{% if fiduciaries.executor.co_primary is defined and fiduciaries.executor.co_primary.full_name %}"
    " and {{ fiduciaries.executor.co_primary.full_name | name }}"
    "{% endif %}, as Personal Representative of this will."
    "{% if fiduciaries.executor.successor is defined and fiduciaries.executor.successor.full_name %}"
    " If for any reason {{ fiduciaries.executor.primary.full_name | name }}"
    "{% if fiduciaries.executor.co_primary is defined and fiduciaries.executor.co_primary.full_name %}"
    " and {{ fiduciaries.executor.co_primary.full_name | name }}"
    "{% endif %} fails or ceases to serve, I appoint "
    "{% if fiduciaries.executor.second_successor is defined and fiduciaries.executor.second_successor.full_name %}"
    "{{ fiduciaries.executor.successor.full_name | name }}"
    "{% if fiduciaries.executor.co_successor is defined and fiduciaries.executor.co_successor.full_name %}"
    " and {{ fiduciaries.executor.co_successor.full_name | name }}"
    "{% endif %}; and {{ fiduciaries.executor.second_successor.full_name | name }}"
    "{% if fiduciaries.executor.co_second_successor is defined and fiduciaries.executor.co_second_successor.full_name %}"
    " and {{ fiduciaries.executor.co_second_successor.full_name | name }}"
    "{% endif %}, in the order named,"
    "{% else %}"
    "{{ fiduciaries.executor.successor.full_name | name }}"
    "{% if fiduciaries.executor.co_successor is defined and fiduciaries.executor.co_successor.full_name %}"
    " and {{ fiduciaries.executor.co_successor.full_name | name }}"
    "{% endif %}"
    "{% endif %} as Personal Representative."
    "{% endif %}"
    " I request that my Personal Representative be appointed Special Personal Representative on application therefor."
)

PR_NOMINATION_PREFIX = "I nominate {{ fiduciaries.executor.primary.full_name | name }}"


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def fix_pr_nomination(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith(PR_NOMINATION_PREFIX):
            replace_paragraph_text(paragraph, PR_NOMINATION_PARAGRAPH)
            return


def remove_legacy_duplicate_paragraphs(document: Document) -> int:
    removed = 0
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if any(text.startswith(prefix) for prefix in DELETE_PARAGRAPH_PREFIXES):
            delete_paragraph(paragraph)
            removed += 1
    return removed


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def patch_family_conditionals_xml(xml: str) -> str:
    """Reuse the family/disinheritance conditional patches from patch_family_template_conditionals."""
    sys.path.insert(0, str(ROOT / "tools"))
    from patch_family_template_conditionals import patch_xml

    return patch_xml(xml)


def apply_family_conditionals(path: Path) -> bool:
    with zipfile.ZipFile(path, "r") as zin:
        original = zin.read("word/document.xml").decode("utf-8")
        patched = patch_family_conditionals_xml(original)
        if patched == original:
            return False

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    data = patched.encode("utf-8")
                zout.writestr(item, data)

    path.write_bytes(out.getvalue())
    return True


def convert_document(document: Document) -> None:
    fix_pr_nomination(document)
    removed = remove_legacy_duplicate_paragraphs(document)
    strip_orphaned_section_references(document)
    if removed:
        print(f"Removed {removed} legacy duplicate PR paragraph(s)")


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    if apply_family_conditionals(TARGET):
        print("Applied family/disinheritance conditional patches")
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
