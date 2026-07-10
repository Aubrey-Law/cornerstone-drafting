#!/usr/bin/env python3
"""Convert the Cornerstone Realty Trust form into a real-estate-transfer template.

Replaces bracket placeholders with docxtemplater `{field}` keys from
`buildTemplateFields()` and writes templates/real-estate-trust.docx.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "source-forms" / "trust" / "cornerstone" / "Realty Trust.docx"
TARGET = ROOT / "templates" / "real-estate-trust.docx"

# Order matters: longer / more specific strings first.
TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    ("[NUMBER OF YEARS OF SUCCESSIVE TRUST TERM]", "{trust_term_extension_years}"),
    ("[PROPERY COUNTY]", "{property_county}"),
    ("[REALTY TRUST NAME]", "{trust_name}"),
    ("[NAME OF TRUSTEE]", "{trustees}"),
    ("[OrdinalDocDate]", "{execution_date}"),
    ("[SIGNING COUNTY]", "{certificate_execution_county}"),
    ("[BENEFICIARY]", "{beneficial_owners}"),
    ("[TRUSTEE]", "{trustees}"),
    ("[DocDate]", "{execution_date}"),
    ("[his/her/their]", "their"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
