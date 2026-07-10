#!/usr/bin/env python3
"""Convert the Cornerstone Delegation of Powers of Parent or Guardian_cv1 form.

Replaces sample client data with Jinja placeholders and writes
templates/temporary-delegation-parental-powers.docx for the
temporary-delegation-parental-powers package.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "source-forms"
    / "guardianship"
    / "cornerstone"
    / "Delegation of Powers of Parent or Guardian_cv1.docx"
)
TARGET = ROOT / "templates" / "temporary-delegation-parental-powers.docx"

PARENT_INTRO_MARRIED = (
    'We, {{ client.full_name | name }} and {{ spouse.full_name | name }}, '
    '("the Parents" "we" "us" "our"), both of {{ client.city }}, {{ client.state }}, state:'
)
PARENT_INTRO_SINGLE = (
    'I, {{ client.full_name | name }}, ("the Parent" "I" "me" "my"), '
    'of {{ client.city }}, {{ client.state }}, state:'
)

WITNESS_ATTESTATION_MARRIED = (
    "On this {{ execution.display_ordinal_doc_date }}, "
    "{{ client.full_name | name }} and {{ spouse.full_name | name }}, "
    "both of {{ client.city }}, in the County of {{ execution.display_signing_county | name }} "
    "and Commonwealth of Massachusetts, signed the foregoing instrument consisting of ______ pages "
    "in our presence, declaring it to be their Massachusetts Temporary Delegation of Powers of Parents, "
    "and as witnesses both of us being over the age of 18 and neither having been appointed Temporary "
    "Agent in this instrument, we sign our names as witnesses."
)

WITNESS_ATTESTATION_SINGLE = (
    "On this {{ execution.display_ordinal_doc_date }}, "
    "{{ client.full_name | name }}, of {{ client.city }}, "
    "in the County of {{ execution.display_signing_county | name }} "
    "and Commonwealth of Massachusetts, signed the foregoing instrument consisting of ______ pages "
    "in our presence, declaring it to be {{ client.pronoun_possessive }} Massachusetts Temporary "
    "Delegation of Powers of Parent, and as witnesses both of us being over the age of 18 and neither "
    "having been appointed Temporary Agent in this instrument, we sign our names as witnesses."
)

NOTARY_ACKNOWLEDGMENT_MARRIED = (
    "On this {{ execution.display_ordinal_doc_date }}, before me, the undersigned notary public "
    "personally appeared {{ client.full_name | name }}, and {{ spouse.full_name | name }}, the Parents, "
    "and _________________________ and _________________________, the Witnesses, who proved to me "
    "through satisfactory evidence, which were Massachusetts Drivers Licenses to be the Parents and "
    "Witnesses, respectively, and whose names are signed to the above instrument, and acknowledged "
    "to me that they signed it voluntarily for its stated purposes."
)

NOTARY_ACKNOWLEDGMENT_SINGLE = (
    "On this {{ execution.display_ordinal_doc_date }}, before me, the undersigned notary public "
    "personally appeared {{ client.full_name | name }}, the Parent, and _________________________ "
    "and _________________________, the Witnesses, who proved to me through satisfactory evidence, "
    "which were Massachusetts Drivers Licenses to be the Parent and Witnesses, respectively, and "
    "whose names are signed to the above instrument, and acknowledged to me that "
    "{{ client.pronoun_subject }} signed it voluntarily for its stated purposes."
)

SIGNATURE_LINE_MARRIED = "{{ client.full_name | name }}\t\t{{ spouse.full_name | name }}"
SIGNATURE_LINE_SINGLE = "{{ client.full_name | name }}"

# Order matters: longer / more specific strings first.
TEXT_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "We, JULIE H. VENABLES and MATTHEW VENABLES, (“the Parents” “we” “us” “our”), both of Needham, Massachusetts, state:",
        "{% if has_spouse %}"
        + PARENT_INTRO_MARRIED
        + "{% else %}"
        + PARENT_INTRO_SINGLE
        + "{% endif %}",
    ),
    ("Name:\tSUSAN HERMAN", "Name:\t{{ guardians.temporary.full_name | name }}"),
    ("Name:\tMICHAEL HERMAN", "Name:\t{{ guardians.temporary_successor.full_name | name }}"),
    ("JULIE H. VENABLES\t\tMATTHEW VENABLES", "{% if has_spouse %}" + SIGNATURE_LINE_MARRIED + "{% else %}" + SIGNATURE_LINE_SINGLE + "{% endif %}"),
    ("SUSAN HERMAN, Temporary Agent", "{{ guardians.temporary.full_name | name }}, Temporary Agent"),
    ("JULIE H. VENABLES", "{{ client.full_name | name }}"),
    ("MATTHEW VENABLES", "{{ spouse.full_name | name }}"),
    ("SUSAN HERMAN", "{{ guardians.temporary.full_name | name }}"),
    ("MICHAEL HERMAN", "{{ guardians.temporary_successor.full_name | name }}"),
    ("this 30th day of July, 2026", "this {{ execution.display_ordinal_doc_date }}"),
    ("On this 30th day of July, 2026", "On this {{ execution.display_ordinal_doc_date }}"),
    ("COUNTY OF NORFOLK", "COUNTY OF {{ execution.display_signing_county | name }}"),
    ("[Notary Commission]", "{{ execution.display_notary_commission }}"),
    ("[UNANSWERED: Notary name]", ""),
    ("Needham, Massachusetts", "{{ client.city }}, {{ client.state }}"),
    ("Needham", "{{ client.city }}"),
]


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def convert_text(text: str) -> str:
    converted = text
    for before, after in TEXT_REPLACEMENTS:
        converted = converted.replace(before, after)
    return converted


def replace_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def convert_paragraph(paragraph: Paragraph) -> None:
    original = paragraph_text(paragraph)
    converted = convert_text(original)
    if converted != original:
        replace_paragraph_text(paragraph, converted)


def convert_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                convert_paragraph(paragraph)


def apply_children_table_fixes(document: Document) -> None:
    if not document.tables:
        return

    table = document.tables[0]
    if len(table.rows) < 2:
        return

    header_cells = [cell.text.strip() for cell in table.rows[0].cells]
    if header_cells != ["Name", "Address", "Date of Birth"]:
        return

    # docxtpl repeats the row between {%tr for %} and {%tr endfor %} tags.
    while len(table.rows) > 2:
        table._tbl.remove(table.rows[2]._tr)

    content_row = table.rows[1]
    content_row.cells[0].paragraphs[0].text = "{{ child.full_name | name }}"
    if len(content_row.cells) > 1:
        content_row.cells[1].paragraphs[0].text = (
            "{{ client.street_address }}, {{ client.city }}, {{ client.state }}"
        )
    if len(content_row.cells) > 2:
        content_row.cells[2].paragraphs[0].text = "{{ child.date_of_birth }}"

    open_tr = deepcopy(content_row._tr)
    content_row._tr.addprevious(open_tr)
    close_tr = deepcopy(content_row._tr)
    content_row._tr.addnext(close_tr)

    open_cells = table.rows[1].cells
    open_cells[0].paragraphs[0].text = "{%tr for child in minor_children %}"
    if len(open_cells) > 1:
        open_cells[1].paragraphs[0].text = ""
    if len(open_cells) > 2:
        open_cells[2].paragraphs[0].text = ""

    close_cells = table.rows[3].cells
    close_cells[0].paragraphs[0].text = "{%tr endfor %}"
    if len(close_cells) > 1:
        close_cells[1].paragraphs[0].text = ""
    if len(close_cells) > 2:
        close_cells[2].paragraphs[0].text = ""


def apply_paragraph_fixes(document: Document) -> None:
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph_text(paragraph)

        if text == "Address:\t__________________________________________________":
            previous = paragraph_text(paragraphs[index - 1]) if index > 0 else ""
            if "guardians.temporary.full_name" in previous:
                replace_paragraph_text(
                    paragraph,
                    "Address:\t{{ guardians.temporary.address | default('__________________________________________________') }}",
                )
            elif "guardians.temporary_successor.full_name" in previous:
                replace_paragraph_text(
                    paragraph,
                    "{% if guardians.temporary_successor.full_name %}"
                    "Address:\t{{ guardians.temporary_successor.address | default('__________________________________________________') }}"
                    "{% endif %}",
                )
            continue

        if text.startswith("On this {{ execution.display_ordinal_doc_date }}, {{ client.full_name | name }} and {{ spouse.full_name | name }}"):
            replace_paragraph_text(
                paragraph,
                "{% if has_spouse %}"
                + WITNESS_ATTESTATION_MARRIED
                + "{% else %}"
                + WITNESS_ATTESTATION_SINGLE
                + "{% endif %}",
            )
            continue

        if text.startswith(
            "On this {{ execution.display_ordinal_doc_date }}, before me, the undersigned notary public personally appeared {{ client.full_name | name }}, and {{ spouse.full_name | name }}"
        ):
            replace_paragraph_text(
                paragraph,
                "{% if has_spouse %}"
                + NOTARY_ACKNOWLEDGMENT_MARRIED
                + "{% else %}"
                + NOTARY_ACKNOWLEDGMENT_SINGLE
                + "{% endif %}",
            )
            continue

        if text == "We appoint as Alternate Temporary Agent:":
            replace_paragraph_text(
                paragraph,
                "{% if guardians.temporary_successor.full_name %}We appoint as Alternate Temporary Agent:{% endif %}",
            )
            continue

        if text == "Name:\t{{ guardians.temporary_successor.full_name | name }}":
            replace_paragraph_text(
                paragraph,
                "{% if guardians.temporary_successor.full_name %}Name:\t{{ guardians.temporary_successor.full_name | name }}{% endif %}",
            )


def strip_orphaned_section_references(document: Document) -> None:
    body = document.element.body
    for sect_pr in body.xpath("./w:sectPr"):
        for tag in ("headerReference", "footerReference"):
            for reference in list(sect_pr.findall(qn(f"w:{tag}"))):
                sect_pr.remove(reference)


def convert_document(document: Document) -> None:
    for paragraph in document.paragraphs:
        convert_paragraph(paragraph)
    for table in document.tables:
        convert_table(table)
    apply_children_table_fixes(document)
    apply_paragraph_fixes(document)
    strip_orphaned_section_references(document)


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source form: {SOURCE}")

    document = Document(SOURCE)
    convert_document(document)
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(f"Wrote {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
