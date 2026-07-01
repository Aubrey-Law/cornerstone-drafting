#!/usr/bin/env python3
"""Flatten Word runs so Nunjucks/Jinja merge tags are contiguous.

Word often splits a single ``{{ field }}`` or ``{% if %}`` across several runs
(e.g. after a spell-check edit), which leaves ``</w:t></w:r><w:t>`` XML in the
middle of a tag. The Node generator renders the raw document XML, so a split tag
becomes a parse error. This script walks every paragraph (body + table cells)
and merges only the runs that together form an unclosed tag, preserving the
formatting of all other runs.

Usage:
    python3 tools/flatten_runs.py <source.docx> <target.docx>
"""

import re
import sys

from docx import Document
from docx.oxml.ns import qn


def has_unclosed(s: str) -> bool:
    s2 = re.sub(r"\{\{.*?\}\}", "", s, flags=re.S)
    s2 = re.sub(r"\{%.*?%\}", "", s2, flags=re.S)
    return ("{{" in s2) or ("{%" in s2) or s2.endswith("{") or s2.endswith("%") or s2.endswith("}")


def run_text(r) -> str:
    return "".join(t.text or "" for t in r.findall(qn("w:t")))


def set_run_text(r, text: str) -> None:
    ts = r.findall(qn("w:t"))
    for extra in ts[1:]:
        r.remove(extra)
    if ts:
        ts[0].text = text
        ts[0].set(qn("xml:space"), "preserve")
    else:
        t = r.makeelement(qn("w:t"), {})
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)


def flatten_paragraph(p) -> None:
    runs = p._p.findall(qn("w:r"))
    items = [(r, run_text(r)) for r in runs]
    i = 0
    while i < len(items):
        elem, text = items[i]
        if has_unclosed(text):
            merged = text
            j = i + 1
            while j < len(items) and has_unclosed(merged):
                merged += items[j][1]
                j += 1
            set_run_text(elem, merged)
            for k in range(i + 1, j):
                items[k][0].getparent().remove(items[k][0])
            i = j
        else:
            i += 1


def all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def main() -> None:
    if len(sys.argv) != 3:
        print(__doc__)
        raise SystemExit(2)
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(src)
    for paragraph in all_paragraphs(doc):
        flatten_paragraph(paragraph)
    doc.save(dst)
    print(f"flattened {src} -> {dst}")


if __name__ == "__main__":
    main()
